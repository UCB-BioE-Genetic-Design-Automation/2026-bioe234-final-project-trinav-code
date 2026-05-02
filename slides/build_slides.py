"""Generate delivery_advisor_slides.pptx for import into Google Slides.

Run:
    python slides/build_slides.py
Then in Google Slides: File -> Import slides -> Upload -> select the .pptx.
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN


# --- Palette ---
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x1A, 0x73, 0xE8)   # blue
GOOD = RGBColor(0x1E, 0x8E, 0x3E)     # green
WARN = RGBColor(0xD9, 0x32, 0x25)     # red
BG_SOFT = RGBColor(0xF3, 0xF6, 0xFA)
BG_CODE = RGBColor(0xF1, 0xF3, 0xF4)
BORDER = RGBColor(0xDA, 0xDC, 0xE0)


def add_text(slide, left, top, width, height, text,
             font_size=14, bold=False, color=INK,
             font_name="Calibri", align=PP_ALIGN.LEFT, line_spacing=1.15):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return box


def add_rich_bullets(slide, left, top, width, height, items,
                     font_size=12, bullet_color=INK, body_color=INK,
                     font_name="Calibri"):
    """items is a list of (bullet_label_or_None, body_text). If bullet is None,
    the row is a plain paragraph."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, (bullet, body) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.25
        if bullet:
            r1 = p.add_run()
            r1.text = "• "
            r1.font.size = Pt(font_size)
            r1.font.bold = True
            r1.font.color.rgb = bullet_color
            r1.font.name = font_name
            r2 = p.add_run()
            r2.text = body
            r2.font.size = Pt(font_size)
            r2.font.color.rgb = body_color
            r2.font.name = font_name
        else:
            r = p.add_run()
            r.text = body
            r.font.size = Pt(font_size)
            r.font.color.rgb = body_color
            r.font.name = font_name
    return box


def add_code_block(slide, left, top, width, height, code, font_size=9):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = BG_CODE
    rect.line.color.rgb = BORDER
    rect.line.width = Pt(0.5)
    rect.shadow.inherit = False
    tf = rect.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(95000)
    tf.margin_right = Emu(95000)
    tf.margin_top = Emu(60000)
    tf.margin_bottom = Emu(60000)
    for i, line in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        r = p.add_run()
        r.text = line if line else " "
        r.font.size = Pt(font_size)
        r.font.name = "Courier New"
        r.font.color.rgb = INK
    return rect


def add_panel(slide, left, top, width, height, fill=BG_SOFT, border=BORDER):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = fill
    rect.line.color.rgb = border
    rect.line.width = Pt(0.5)
    rect.shadow.inherit = False
    rect.text_frame.margin_left = Emu(0)
    rect.text_frame.margin_right = Emu(0)
    rect.text_frame.margin_top = Emu(0)
    rect.text_frame.margin_bottom = Emu(0)
    return rect


def add_arrow_down(slide, center_x, top, height=Inches(0.22)):
    w = Inches(0.14)
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, center_x - w / 2, top, w, height
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = MUTED
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    return arrow


# --- Build ---

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank = prs.slide_layouts[6]

# =====================================================================
# SLIDE 1 — What & How You Use It
# =====================================================================
slide = prs.slides.add_slide(blank)

# Title
add_text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.6),
         "Delivery Strategy Advisor", font_size=32, bold=True, color=INK)

# Subtitle
add_text(
    slide, Inches(0.5), Inches(0.95), Inches(12.3), Inches(0.6),
    "Given a cell type + Cas variant + edit type (or the upstream construct), ranks CRISPR delivery "
    "methods with rationale, KB citations, and a sanity-check on the builder's kit choice.",
    font_size=14, color=MUTED,
)

# LEFT column — Example MCP prompts
left_x = Inches(0.5)
left_y = Inches(1.8)
left_w = Inches(5.4)

add_text(slide, left_x, left_y, left_w, Inches(0.35),
         "Example MCP prompts", font_size=15, bold=True, color=ACCENT)

add_rich_bullets(
    slide, left_x, left_y + Inches(0.4), left_w, Inches(3.0),
    [
        (True, "\"What delivery method should I use for a SpCas9 knockout in HEK293 cells?\""),
        (True, "\"I want to use ABE8e to do a base edit in hepatocytes. What should I use?\"   (cargo-limit edge case)"),
        (True, "\"Here is the construction file from the builder: {…}. What delivery should I use, and does the kit choice make sense?\"   (pipeline integration)"),
    ],
    font_size=12,
)

# Callout panel — the unique angle
callout = add_panel(slide, left_x, Inches(5.3), left_w, Inches(1.5),
                    fill=RGBColor(0xEA, 0xF3, 0xFD), border=ACCENT)
add_text(slide, left_x + Inches(0.2), Inches(5.45), left_w - Inches(0.4), Inches(0.4),
         "What's unusual", font_size=12, bold=True, color=ACCENT)
add_text(
    slide, left_x + Inches(0.2), Inches(5.85), left_w - Inches(0.4), Inches(1.0),
    "The advisor doesn't just recommend a method — it audits the upstream builder's\n"
    "kit_choice and surfaces agreement or disagreement in a kit_choice_consistency block.",
    font_size=12, color=INK,
)

# RIGHT column — Input / Output
right_x = Inches(6.2)
right_w = Inches(6.6)

add_text(slide, right_x, left_y, right_w, Inches(0.35),
         "Input  (from Construction File Builder)", font_size=12, bold=True, color=ACCENT)

input_code = """{
  "cell_type": "HEK293",
  "cas_variant": "SpCas9",
  "edit_type": "knockout",
  "kit_choice": "lentiCRISPRv2",
  "off_target_risk_profile": {
    "total_off_targets": 2, ...
  }
}"""
add_code_block(slide, right_x, left_y + Inches(0.3), right_w, Inches(1.9), input_code, font_size=10)

add_text(slide, right_x, left_y + Inches(2.35), right_w, Inches(0.35),
         "Output  (abbreviated)", font_size=12, bold=True, color=ACCENT)

output_code = """{
  "top_recommendation": "Lipofection",
  "ranked_methods": [
    { "method": "Lipofection",        "suitability": "high",
      "citations": ["lipofection"] },
    { "method": "RNP electroporation", "suitability": "high",
      "citations": ["rnp_electroporation"] }
  ],
  "kit_choice_consistency": {
    "kit_choice":  "lentiCRISPRv2",
    "consistent":  false,
    "note": "Integrating lentivirus adds off-target risk;
             transient lipofection preferred for HEK293 KO."
  }
}"""
add_code_block(slide, right_x, left_y + Inches(2.65), right_w, Inches(4.15), output_code, font_size=10)


# =====================================================================
# SLIDE 2 — Architecture & Testing
# =====================================================================
slide = prs.slides.add_slide(blank)

add_text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.6),
         "Architecture — Hybrid Prefilter + KB-Grounded LLM",
         font_size=28, bold=True, color=INK)

# --- Top: data-flow diagram ---
# Horizontal flow: Input -> prefilter -> render KB -> Claude -> validate -> Output
flow_y = Inches(1.1)
node_h = Inches(0.7)
node_w = Inches(1.9)
gap = Inches(0.25)
start_x = Inches(0.5)


def add_flow_node(x, y, w, h, title, subtitle=None, fill=BG_SOFT, border=BORDER,
                  title_color=INK, subtitle_color=MUTED):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    rect.fill.solid()
    rect.fill.fore_color.rgb = fill
    rect.line.color.rgb = border
    rect.line.width = Pt(0.75)
    rect.shadow.inherit = False
    tf = rect.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(50000)
    tf.margin_right = Emu(50000)
    tf.margin_top = Emu(40000)
    tf.margin_bottom = Emu(40000)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = title_color
    r.font.name = "Calibri"
    if subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.line_spacing = 1.05
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(8.5)
        r2.font.color.rgb = subtitle_color
        r2.font.name = "Calibri"
    return rect


def add_right_arrow(x, y_center, length=Inches(0.22)):
    h = Inches(0.14)
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, x, y_center - h / 2, length, h
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = MUTED
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    return arrow


nodes = [
    ("Input", "cell_type, cas_variant,\nedit_type  /  construct", BG_SOFT, BORDER),
    ("_prefilter", "AAV limit, plant/rodent,\nzygote-only rules", RGBColor(0xFE, 0xF3, 0xE2), RGBColor(0xF2, 0xB1, 0x4F)),
    ("Render KB", "methods.json — 17 methods,\ncargo table, heuristics", BG_SOFT, BORDER),
    ("Claude Haiku 4.5", "KB + schema in system,\ncase + kit_choice in user", RGBColor(0xEA, 0xF3, 0xFD), ACCENT),
    ("Validate citations", "every rec cites a real\nKB key", RGBColor(0xE6, 0xF4, 0xEA), GOOD),
    ("Output", "top_rec · ranked methods\nw/ citations · kit_consistency", BG_SOFT, BORDER),
]

x = start_x
y_center = flow_y + node_h / 2
for i, (title, sub, fill, border) in enumerate(nodes):
    add_flow_node(x, flow_y, node_w, node_h, title, sub, fill=fill, border=border)
    if i < len(nodes) - 1:
        add_right_arrow(x + node_w + Inches(0.02), y_center, length=gap - Inches(0.04))
    x += node_w + gap

# --- Bottom: two columns ---
bottom_y = Inches(2.2)

# LEFT: What makes it unique
lx = Inches(0.5)
lw = Inches(6.2)
add_text(slide, lx, bottom_y, lw, Inches(0.4),
         "What makes it unique", font_size=16, bold=True, color=ACCENT)

unique_items = [
    (True, "Deterministic prefilter kills infeasible methods before the LLM sees them — no hallucinated \"single AAV for ABE8e.\""),
    (True, "KB lives as versioned JSON (data/methods.json) — every rule change is reviewed in git diff."),
    (True, "Citation enforcement — each recommendation cites a real KB key; unknown keys trigger _citation_warnings."),
    (True, "Pipeline audit — when called with a construct, compares the upstream kit_choice to the top recommendation and explains any disagreement."),
]
add_rich_bullets(slide, lx, bottom_y + Inches(0.5), lw, Inches(3.8), unique_items,
                 font_size=12, bullet_color=ACCENT)

# RIGHT: Testing
rx = Inches(6.95)
rw = Inches(5.85)
add_text(slide, rx, bottom_y, rw, Inches(0.4),
         "Testing", font_size=16, bold=True, color=GOOD)

# Pytest panel
add_panel(slide, rx, bottom_y + Inches(0.5), rw, Inches(2.25),
          fill=RGBColor(0xE6, 0xF4, 0xEA), border=GOOD)
add_text(slide, rx + Inches(0.25), bottom_y + Inches(0.6), rw - Inches(0.5), Inches(0.35),
         "26 pytests  (tests/test_delivery_advisor.py)",
         font_size=12, bold=True, color=GOOD)
add_rich_bullets(
    slide, rx + Inches(0.25), bottom_y + Inches(0.95), rw - Inches(0.5), Inches(1.8),
    [
        (True, "Input validation, response schema, citation validity"),
        (True, "Biology edge cases: ABE8e > AAV limit, plant cells, in-vivo brain (CNS tropism), clinical T cells"),
        (True, "5 pipeline-integration tests — construct path + kit-choice audit"),
    ],
    font_size=11, bullet_color=GOOD,
)

# Eval harness panel
add_panel(slide, rx, bottom_y + Inches(2.9), rw, Inches(1.55),
          fill=RGBColor(0xEA, 0xF3, 0xFD), border=ACCENT)
add_text(slide, rx + Inches(0.25), bottom_y + Inches(3.0), rw - Inches(0.5), Inches(0.35),
         "20-case accuracy eval  (tests/eval_delivery_advisor.py)",
         font_size=12, bold=True, color=ACCENT)
add_rich_bullets(
    slide, rx + Inches(0.25), bottom_y + Inches(3.35), rw - Inches(0.5), Inches(1.15),
    [
        (True, "Reports top-rec accuracy, citation validity, kit consistency"),
        (True, "Live example: pipeline-kit-mismatch → top=\"lipofection\", consistent=False ✓"),
    ],
    font_size=11, bullet_color=ACCENT,
)


# --- Save ---
out = Path(__file__).parent / "delivery_advisor_slides.pptx"
prs.save(out)
print(f"Wrote {out}")
